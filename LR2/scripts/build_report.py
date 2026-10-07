#!/usr/bin/env python3
"""Собирает отчёт Word и его текстовую копию из проверенных артефактов."""
from pathlib import Path
import csv,json,re,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
V=ROOT/'artifacts/v03'
OUT=ROOT/'reports/ЛР02_Вариант_3_Отчёт.docx'
def read(name): return (V/name).read_text(encoding='utf-8').strip()
def body(name): return read(name).split('\n',1)[1].strip()
env=json.loads((ROOT/'reports/environment.json').read_text())
sections=[]
sections.append(f'''# 1 Цель сценарий данные и среда

Цель работы — проверить происхождение и условия использования данных, составить паспорт набора, карточку генератора и реестр рисков. По варианту 3 выполнен аудит 12 синтетических записей: найдены 3 ошибки ND и 1 предупреждение о дубликате. По всем находкам записаны решения. После исключения 3 ND-записей и карантина item10 выборка из 8 записей даёт 0 ошибок и 0 предупреждений по учебным правилам. Реестр 7 рисков прошёл проверку RESULT: PASS.

## Задание варианта 3

Сценарий — афиша фестиваля с печатью и продажей билетов; профиль — NoDerivatives. Номер определяется как (1−1)×6+3=3. Политика сохранена без изменения: commercial_use=true, redistribution=true, derivatives=true, share_alike_acceptable=false. Для реестра обязательны Intellectual Property и Information Integrity.

Критерии выполнения: ожидания до аудита; решения для 100 % находок, включая предупреждения; паспорт с 7 и карточка с 9 заполненными разделами; не менее 5 рисков и 4 категорий; проверка PASS; разделение инженерных и юридических вопросов.

## Предварительные ожидания

07.10.2026 в 19:38:09 MSK до первого аудита сохранён reports/expectations_before_audit.md. По коду генератора ожидались 3 ошибки R4 для ND и одно фоновое предупреждение R8, R9 либо R10; R3, R5, R6 и R7 не ожидались. Конкретные объекты не прогнозировались. После запуска подтверждены три R4 и одно R9; остальные ожидаемые отсутствия срабатываний также подтверждены.

## Происхождение данных

Вход создан штатным lab02_make_variant.py для варианта 3. Файлы медиа отсутствуют: имена, источники example.org, даты и правообладатели условны. SHA-256 — хеши строк, один из которых продублирован намеренно. Работа анализирует метаданные и политику; права на реальные материалы таким способом не подтверждаются. Результаты ЛР 1 взяты из собственного сохранённого паспорта варианта 3 от 24.09.2026.

## Среда выполнения ЛР 2

Фактически использованы {env['implementation']} {env['python']}, {env['platform']}. Пять учебных скриптов используют стандартную библиотеку. Для оформления отчёта применён python-docx из рабочего runtime; это не зависимость аудита. Версии, команды, коды завершения и выводы сохранены в reports/environment.json и reports/run_log.json. Версии среды ЛР 1 приведены отдельно в карточке модели и не подменяются текущими.
''')

audit=body('audit_report.md')
sections.append('# 2 Результаты аудита и решения\n\n'+audit.split('## item03')[0] + '''
## Решения по объектам ND

item03 — изображение, item04 — аудио, item11 — текст. Все три записи CC-BY-ND-4.0 исключены из выборки для публикуемых производных. Для item03 предполагается замена собственным процедурным фоном либо материалом CC0 или CC-BY с подтверждёнными правами; для item11 — собственным проверенным текстом. Аудио item04 для статической афиши не требуется; для будущего проморолика необходима подходящая замена или разрешение.

Альтернатива исключению — отдельное документированное разрешение правообладателя на адаптацию и коммерческую публикацию. Разрешение не запрашивалось и не получено; лицензии в manifest.csv не менялись. Использование исходного материала без адаптации возможно лишь как отдельно проверенная ветка макета с сохранением атрибуции; в данной работе она не реализована.

## Решение по предупреждению R9

item10 помещён в карантин: его SHA-256 совпадает с item01, но у item10 указаны Apache-2.0 и Автор-09, а у item01 — MIT и Автор-15. Совпадение хеша не позволяет выбрать верную лицензию и не подтверждает права. item01 сохранён только для учебной проверки. В реальном проекте обе записи пары блокируются до сравнения файлов и подтверждения происхождения. Подробные решения по всем четырём находкам сохранены в audit_decisions.json.

## Проверка после решений

В manifest_selected.csv оставлены item01, item02, item05, item06, item07, item08, item09, item12. Повторный аудит при исходной policy.json: объектов 8; ошибок 0; предупреждений 0. Исходный manifest.csv и первоначальные результаты аудита сохранены. Отсутствие находок означает соответствие записанных метаданных учебным правилам. Это не юридическая консультация и не разрешение реального выпуска.
''')

data=body('data_card.md')
cut=data.index('## 5 Uses')
sections.append('# 3 Паспорт набора данных\n\n'+data[:cut].strip())
sections.append('# 3 Паспорт набора данных продолжение\n\n'+data[cut:].strip()+'''

## Разделение объектов прав

Лицензии входных материалов записаны отдельно в манифесте. Лицензия кода генератора неизвестна и рассматривается в карточке модели. Весов нет, поэтому лицензия весов неприменима. Права на результат генерации и макет афиши не выводятся из лицензии манифеста: требуется проверить права на каждый компонент макета и условия использования кода. Эти три сущности не объединяются в одну лицензию.
''')

model=body('model_card.md')
cut=model.index('## 5 Evaluation')
sections.append('# 4 Карточка модели\n\n'+model[:cut].strip())
sections.append('# 4 Карточка модели продолжение\n\n'+model[cut:].strip())

risks=list(csv.DictReader((V/'risk_register.csv').open(encoding='utf-8',newline='')))
def details(r):
    return f"## {r['risk_id']} {r['category']}\n\n{r['description']}. L={r['likelihood']}, I={r['impact']}, R={int(r['likelihood'])*int(r['impact'])}.\n\nМера: {r['measure']}. Ответственный: {r['owner']}.\n\nПризнак наступления: {r['trigger']}.\n\n"
summary='| Риск | Категория NIST AI 600 1 | L | I | R |\n|---|---|---|---|---|\n'
for r in sorted(risks,key=lambda r:-int(r['likelihood'])*int(r['impact'])):
    summary+=f"| {r['risk_id']} | {r['category']} | {r['likelihood']} | {r['impact']} | {int(r['likelihood'])*int(r['impact'])} |\n"
sections.append('# 5 Реестр рисков\n\n'+body('risk_context.md')+'\n\n'+summary+'\n\n'+details(risks[0]))
sections.append('# 5 Реестр рисков продолжение\n\n'+''.join(details(r) for r in risks[1:4]))
sections.append('# 5 Реестр рисков продолжение\n\n'+''.join(details(r) for r in risks[4:])+'''
## Результат проверки

python src/lab02_risks.py --csv artifacts/v03/risk_register.csv --out artifacts/v03/risk_ranking.md завершился с кодом 0 и RESULT: PASS. Категории совпадают с перечнем скрипта. Дополнительно проверены уникальность идентификаторов и наличие обязательных Intellectual Property и Information Integrity; штатный валидатор обязательность категорий конкретного варианта сам не проверяет. Проверка не подтверждает, что оценки статистически измерены или меры реализованы в реальном проекте.
''')

sections.append('''# 6 Анализ ограничений и спорных решений

## Граница автоматического аудита

Скрипт видит строки CSV, но не открывает источник, медиафайл, лицензионное уведомление или согласие. Он не устанавливает полномочия правообладателя, наличие людей в содержимом, достаточность согласия, срок или территорию разрешения. Поэтому корректность записи лицензии и права на конкретный объект требуют первоисточника и документов. Хеши здесь условны; R9 нельзя трактовать как доказанную идентичность реальных файлов.

Для реального проекта требуются оригиналы файлов и их хеши, сохранённая страница или документ лицензии с датой проверки, подтверждение полномочий автора, сведения об атрибуции и договор или разрешение на нужный способ использования. Для пары item01/item10 нужно установить, почему различаются лицензия и правообладатель. Самостоятельное удаление одной строки не решает правовой конфликт.

## Упрощение правила ND

Сообщение R4 формулируется как запрет производных. В CC BY-ND 4.0 раздел 2(a)(1)(B) допускает создание и воспроизведение Adapted Material, но не его Share. В нашем сценарии присутствует именно публикация производной афиши, поэтому конфликт сохраняется. Для внутреннего непубликуемого прототипа blanket-срабатывание R4 может переоценивать запрет и требует отдельного разбора [2].

Раздел 2(a)(4) разрешает необходимые технические изменения формата, которые сами по себе не создают Adapted Material [2]. Поэтому обычная смена формата или техническое сжатие не должны автоматически приравниваться к адаптации. Кадрирование, коллаж, изменение художественного содержания и способ включения в макет оцениваются по обстоятельствам и применимому праву. Простое указание автора не снимает запрет распространения адаптации.

## Какие вопросы решены инженерно

Инженерно выполнены создание набора, аудит, сохранение исходных метаданных, исключение ND-записей, карантин item10, повторная проверка и структурная валидация рисков. Юридического участия требуют квалификация конкретного макета как адаптации, достаточность отдельного разрешения, права реального автора и основания коммерческого использования кода. Изменение сценария на некоммерческий не снимает ND при публикации адаптаций.

## Предложение по улучшению правил

Для R4 следует разделить создание адаптации и её распространение: при ND, derivatives=true и redistribution=true выдавать конфликт публикации; без передачи адаптации — информационную находку с проверкой иных оснований. Для R9 нужно отдельно отмечать несовпадение лицензионных метаданных дубликатов и блокировать обе записи до проверки. Полезны также проверка уникальности id, целостности полей, источника лицензии и обязательных категорий варианта. В предоставленные скрипты изменения не вносились, чтобы сохранить учебную методику.
''')

sections.append('''# 7 Вывод о пригодности набора

Исходный набор из 12 записей не соответствует заданной политике производных публикуемых материалов: три объекта CC-BY-ND-4.0 дают ошибки R4. Предупреждение R9 по item10 также требует решения из-за противоречивых сведений о правах дубликата. ND-записи исключены, item10 помещён в карантин, а исходные сведения сохранены.

Для учебного сценария выборка из 8 записей проходит повторный аудит без ошибок и предупреждений. Политика не менялась, новые лицензии, даты или разрешения не придуманы. Установлена пообъектная процедура допуска. Однако для реального производства требуется подтвердить права по всем первоисточникам; отдельная проверка нужна для item01 и item10, даже если одна запись убрана из выборки.

Паспорт данных содержит 7 заполненных разделов; карточка процедурного генератора — 9. Для генератора нет обучения и весов, а отсутствие лицензии кода задокументировано как ограничение. Ответственный за эксперимент и документацию указан, без присвоения авторства учебного кода. Метрики ЛР 1 перенесены из сохранённого паспорта и не представлены как новый эксперимент.

Реестр содержит 7 рисков из 7 категорий NIST, включая обязательные Intellectual Property и Information Integrity. Проверка дала RESULT: PASS. Наибольшая оценка у риска публикации ND-адаптации — 20; далее риск ошибочного происхождения дубликата — 16. Для каждого риска указаны мера, роль ответственного и наблюдаемый признак.

Цель лабораторной выполнена в пределах учебной модели. Документы и отсутствие автоматических находок не подтверждают юридическую чистоту реального набора. Работа не является юридической консультацией; решение о реальном выпуске принимает ответственный с участием специалиста по правам.

## Происхождение кода и выполненные действия

Пять скриптов ЛР 2 предоставлены методическим комплектом и сохранены без изменения. В ходе выполнения создан вариант 3, заранее записаны ожидания, проведены два аудита, оформлены решения, две карточки и реестр, выполнена проверка рисков, подготовлены отчёт и архив. Реальные медиаматериалы не собирались, афиша не публиковалась, запросы правообладателям не отправлялись.

## Комплект для сдачи

Отчёт находится в reports/ЛР02_Вариант_3_Отчёт.docx. Архив ЛР02_Вариант_3_Сдача.zip содержит папку LR2 с исходными скриптами, входным паспортом ЛР 1, манифестами, политикой, исходным аудитом, решениями, паспортом данных, карточкой модели, реестром рисков, журналом и проверками. Карта файлов и команда повторной проверки приведены в README.md. Архив находится рядом с папкой LR2 в проекте.
''')

sections.append('''# 8 Ход работы и воспроизведение

Команды ниже выполняются из папки LR2. Требуется Python 3.10 или новее. Учебные вычисления не используют сторонних пакетов. При повторном выполнении ожидания первоначального запуска не заменяются и не датируются заново.

```text
python src/lab02_make_variant.py --variant 3 --out artifacts/v03
python src/lab02_audit.py --manifest artifacts/v03/manifest.csv --policy artifacts/v03/policy.json --out artifacts/v03
python src/lab02_model_card.py --passport inputs/lr01_variant03_passport.json --owner "Власюк Данил Витальевич" --out artifacts/v03/model_card_draft.md
python src/lab02_risks.py --template artifacts/v03/risk_register_template.csv
python scripts/prepare_results.py
python src/lab02_audit.py --manifest artifacts/v03/manifest_selected.csv --policy artifacts/v03/policy.json --out artifacts/v03/selected_audit
python src/lab02_risks.py --csv artifacts/v03/risk_register.csv --out artifacts/v03/risk_ranking.md
python scripts/verify_submission.py
```

Карточка из штатного скрипта служит заготовкой; окончательная model_card.md дополнена определениями метрик, двумя условиями эксперимента, ответственным и ограничением лицензии кода. data_card.md — итоговый паспорт данных; автоматически созданные заготовки сохранены как свидетельство работы скриптов.

Исходные выводы и коды завершения сохранены в reports/run_log.json; полный вывод проверки реестра — в reports/risk_validation.log. SHA-256 скриптов и входов находится в reports/source_sha256.json. Скрипт verify_submission.py создаёт вариант заново во временной папке, сравнивает данные, пересчитывает аудит и проверяет покрытия решений и категорий. Отчёт Word при этом не изменяется.

## Основные файлы артефактов

| Назначение | Файл относительно LR2 |
|---|---|
| Прогноз до первого аудита | reports/expectations_before_audit.md |
| Вход и политика | artifacts/v03/manifest.csv и policy.json |
| Аудит с решениями | artifacts/v03/audit_report.md и audit_decisions.json |
| Выборка и её аудит | artifacts/v03/manifest_selected.csv и selected_audit/audit_report.json |
| Паспорт набора | artifacts/v03/data_card.md |
| Карточка модели | artifacts/v03/model_card.md |
| Реестр и ранжирование | artifacts/v03/risk_register.csv и risk_ranking.md |
| Проверка и журнал | reports/verification.json и run_log.json |
''')

sections.append('''# 9 Ответы на контрольные вопросы

## 1 Три разных объекта прав

Лицензия входных данных определяет способы использования материалов; лицензия весов — способы использования и распространения параметров обученной модели; права на результат зависят от условий инструмента, состава результата и применимого права. Ошибка — считать результат свободным только потому, что вход был открытым. Здесь весов нет, а лицензия кода неизвестна.

## 2 Условия Creative Commons

BY требует предусмотренной атрибуции; SA — совместимых условий распространения адаптации; NC ограничивает коммерческое использование; ND не разрешает распространение адаптации. Для коммерческой афиши NC создаёт конфликт с коммерческим сценарием, ND — с публикацией адаптации. SA в заданной учебной политике также требует решения [2, 3].

## 3 Нестандартная запись лицензии

«CC BY 4» не является точным идентификатором из учебного каталога. Без первоисточника нельзя исключить ошибку или другое условие. Стандартная запись — CC-BY-4.0. Нужно запросить подтверждение лицензии, а не автоматически переименовать запись.

## 4 Две находки по одному объекту

В демонстрационном варианте CC-BY-NC-SA-4.0 нарушает коммерческое условие R3 и вызывает предупреждение R5 при неприемлемом ShareAlike. Исключение объекта снимает обе проблемы для выборки. В нашем варианте основной конфликт — R4, поэтому эта пара правил не срабатывает.

## 5 Ошибка и предупреждение

Ошибка фиксирует конфликт с правилами учебной политики, предупреждение — проблему, требующую решения и проверки. R9 по item10 нельзя оставить без решения: совпадающий хеш сопровождается разными сведениями о правах. Выбран карантин.

## 6 Согласия при публикации

Для четырёх записей с персональными данными, где одна без согласия, две с документом и одна с n-a: при публикации 2 ошибки R6 и 2 предупреждения R7; без публикации 2 ошибки R6 и 0 предупреждений R7. Это учитывает только указанные дефекты.

## 7 Расчёт риска

При L=3 и I=5 оценка R=15. Пара L=4, I=3 даёт 12, а L=2, I=5 — 10; выше ранжируется первая. Равенство произведений не означает одинаковую природу последствий.

## 8 Пробелы паспорта ЛР 1

Из паспорта нельзя вывести лицензию чужого кода и ответственного. В карточке ответственным указан исполнитель лабораторной работы; лицензия остаётся неустановленной до ответа автора или преподавателя. Это показывает, что технических параметров недостаточно для документации ответственности и прав.

## 9 Наблюдаемый признак

Без признака нельзя установить, что риск наступил, и запустить меру. Для Data Privacy наблюдаемое событие — обнаружена запись personal_data=yes без документа-согласия либо документ не покрывает публикацию. Формулировка «возможна утечка» событие не задаёт.

## 10 Пределы скрипта

Аудит не подтверждает права, согласия и подлинность содержимого. Нужны оригинал файла, первоисточник лицензии, сведения об авторе и основаниях разрешения, а при спорной адаптации — правовая оценка конкретного макета. Отсутствие ошибок по CSV не равно проверке этих документов.
''')

questions=sections.pop()
qcut=questions.index('## 6 Согласия')
sections.append(questions[:qcut].strip())
sections.append('# 9 Ответы на контрольные вопросы продолжение\n\n'+questions[qcut:].strip())

sections.append('''# Источники

[1] Методические указания «Лабораторная работа № 2. Карточка модели, лицензии, происхождение данных и анализ рисков», предоставленные DOCX и ZIP; разделы 3–5, вариант 3 в разделе 7, требования к отчёту в разделе 8. Пять учебных скриптов — тот же комплект.

[2] Creative Commons. Attribution NoDerivatives 4.0 International. Юридический текст, разделы 2(a)(1)(B), 2(a)(4), 3(a). Проверено 07.10.2026 в Dia. https://creativecommons.org/licenses/by-nd/4.0/legalcode.en

[3] Creative Commons. Attribution 4.0 International. Условия атрибуции и изменений; проверено 07.10.2026. https://creativecommons.org/licenses/by/4.0/

[4] Apache Software Foundation. Apache License Version 2.0. Раздел 4; проверено 07.10.2026. https://www.apache.org/licenses/LICENSE-2.0

[5] SPDX. MIT License. Текст лицензии и идентификатор; проверено 07.10.2026. https://spdx.org/licenses/MIT.html

[6] NIST AI 600-1. Artificial Intelligence Risk Management Framework Generative Artificial Intelligence Profile, 2024. Перечень категорий взят из предоставленных методических указаний и lab02_common.py. Первичный документ: https://doi.org/10.6028/NIST.AI.600-1

[7] Gebru T. и др. Datasheets for Datasets. arXiv:1803.09010, 2018. Семь разделов паспорта применены по шаблону методических указаний. https://arxiv.org/abs/1803.09010

[8] Mitchell M. и др. Model Cards for Model Reporting. arXiv:1810.03993, 2018. Девять разделов карточки применены по шаблону методических указаний. https://arxiv.org/abs/1810.03993

[9] Паспорт и результаты ЛР 1, вариант 3, 24.09.2026. Входные копии: inputs/lr01_variant03_passport.json и inputs/lr01_variant03_results.json; исходный генератор и скрипт эксперимента — inputs/lr01_src/. Числовые данные карточки модели взяты из этого паспорта.
''')

title='Лабораторная работа 2\nКарточка модели лицензии происхождение данных и анализ рисков'
def section_body(text):
    return text.strip().split('\n',1)[1].strip()

def demote(text, levels=1):
    return '\n'.join('#'*levels+line if line.startswith('#') else line for line in text.splitlines())

task=section_body(sections[0])
origin_pos=task.index('## Происхождение данных')
env_pos=task.index('## Среда выполнения ЛР 2')
purpose=task[:origin_pos].strip()
origin=task[origin_pos:env_pos].split('\n',1)[1].strip()
environment=task[env_pos:].split('\n',1)[1].strip()
conclusion=section_body(sections[10])
contribution_pos=conclusion.index('## Происхождение кода и выполненные действия')
contribution=conclusion[contribution_pos:].strip()
conclusion=conclusion[:contribution_pos].strip()
explanation_of_rights=section_body(sections[3]).split('## Разделение объектов прав',1)[1].strip()
parts=[
    '## 1. Цель, постановка задачи и критерии успешности\n\n'+demote(purpose),
    '## 2. Происхождение данных, ограничения и лицензии\n\n'+origin+'\n\n'+explanation_of_rights,
    '## 3. Среда, версии и исходные скрипты\n\n'+environment,
    '## 4. Ход работы и команды запуска\n\n'+demote(section_body(sections[11])),
    '## 5. Результаты аудита, карточки и реестр рисков\n\n'+
        '### 5.1. Аудит и решения по всем находкам\n\n'+demote(section_body(sections[1]),2)+'\n\n'+
        '### 5.2. Паспорт набора данных\n\n'+demote(section_body(sections[2]),2)+'\n\n'+demote(section_body(sections[3]).split('## Разделение объектов прав',1)[0].strip(),2)+'\n\n'+
        '### 5.3. Карточка генератора из ЛР 1\n\n'+demote(section_body(sections[4]),2)+'\n\n'+demote(section_body(sections[5]),2)+'\n\n'+
        '### 5.4. Реестр рисков и результат проверки\n\n'+demote(section_body(sections[6]),2)+'\n\n'+demote(section_body(sections[7]),2)+'\n\n'+demote(section_body(sections[8]),2),
    '## 6. Диагностика, ограничения и спорные решения\n\n'+demote(section_body(sections[9])),
    '## 7. Вывод по цели\n\n'+conclusion,
    '## 8. Архив, репозиторий и выполненные действия\n\n'+demote(contribution),
    '## 9. Ответы на контрольные вопросы\n\n'+demote(section_body(sections[12]))+'\n\n'+demote(section_body(sections[13])),
    '## Источники\n\n'+section_body(sections[14]),
]
markdown='''# Лабораторная работа № 2
## Карточка модели, лицензии, происхождение данных и анализ рисков

**Дисциплина:** Искусственный интеллект в креативных технологиях

**Студент:** Власюк Данил Витальевич

**Группа:** Т26ИИКТ-МИК11

**Вариант:** 3 — коммерческая афиша фестиваля, профиль NoDerivatives

**Дата выполнения:** 07.10.2026

**Репозиторий:** https://github.com/mart1ny/AI_in_creative_media.git

**Папка работы:** `LR2/`

Выполнен аудит синтетического набора из 12 записей. Найдены три ошибки R4 и одно предупреждение R9; по каждой находке сохранены решение и обоснование. Выборка из восьми записей проходит повторный аудит без находок. Паспорт данных и карточка генератора заполнены; реестр из семи рисков даёт **RESULT: PASS**. Скрипты, исходные данные, журналы и результаты сохранены в репозитории. Это проверка учебной политики по записанным метаданным, а не подтверждение прав на реальные материалы.

---

'''+ '\n\n---\n\n'.join(parts)+'\n'
(ROOT/'reports/ЛР02_Вариант_3_Отчёт.md').write_text(markdown,encoding='utf-8')
if '--markdown-only' in sys.argv:
    print(ROOT/'reports/ЛР02_Вариант_3_Отчёт.md')
    raise SystemExit(0)

doc=Document()
s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.3);s.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    fonts=st.element.get_or_add_rPr().rFonts
    for attr in list(fonts.attrib):
        if 'Theme' in attr:del fonts.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
normal=doc.styles['Normal'];normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.10
normal.paragraph_format.space_after=Pt(6)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12.5)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.keep_with_next=True
doc.core_properties.title=title.replace('\n',' ')
doc.core_properties.author='Власюк Данил Витальевич'
doc.core_properties.subject='ЛР 2 вариант 3'
p=doc.add_paragraph('Искусственный интеллект в креативных технологиях')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(70)
p=doc.add_paragraph(title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Отчёт по варианту 3');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(55)
for text in ['Выполнил Власюк Данил Витальевич','Группа Т26ИИКТ-МИК11','Сценарий коммерческая афиша фестиваля','Профиль лицензии NoDerivatives']:
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('07 октября 2026 года');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(80)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
s.different_first_page_header_footer=True

def add_table(lines):
    values=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines]
    values=[values[0]]+values[2:]
    n=len(values[0]);t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    width=16.9
    if n==5: widths=[1.3,10.0,1.8,1.8,2.0] if values[0][0]=='Риск' else [0.7,1.4,2.4,1.7,10.7]
    elif n==7: widths=[3.0,2.1,2.1,2.1,2.1,2.1,3.4]
    else: widths=[6,10.9]
    widths=[w*width/sum(widths) for w in widths]
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    for i,row in enumerate(values):
        cells=t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (n==7 and j>0) or (n==5 and (j==0 or values[0][0]=='Риск' and j>1)) else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(text);r.font.size=Pt(9.5 if n==7 else 10.5);r.bold=i==0
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF3');tcpr.append(shading)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def render_md(text):
    lines=text.splitlines();i=0;code=False
    while i<len(lines):
        l=lines[i]
        if l.startswith('```'):code=not code;i+=1;continue
        if not l.strip():i+=1;continue
        if l.startswith('|'):
            group=[]
            while i<len(lines) and lines[i].startswith('|'):group.append(lines[i]);i+=1
            add_table(group);continue
        if l.startswith('# '):doc.add_paragraph(l[2:],'Heading 1')
        elif l.startswith('## '):doc.add_paragraph(l[3:],'Heading 2')
        else:
            p=doc.add_paragraph(l)
            if code:
                p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.name='Courier New';r.font.size=Pt(8.5)
            else:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        i+=1

for section in sections:
    start=len(doc.paragraphs)
    render_md(section.strip())
    doc.paragraphs[start].paragraph_format.page_break_before=True
doc.save(OUT)
print(OUT)
