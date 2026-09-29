import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'reports/LR01_1_report.md'
OUT = ROOT / 'reports/ЛР01_Вариант_3_Отчёт.docx'
RESULT = json.loads((ROOT / 'artifacts/variant03/results.json').read_text())
DIAG = json.loads((ROOT / 'reports/seed_diagnostic.json').read_text())


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)


def add_table(doc, rows, widths=None, small=False):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for ri, items in enumerate(rows):
        cells = table.add_row().cells
        for ci, val in enumerate(items):
            c = cells[ci]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.text = str(val)
            set_cell_border(c)
            if widths:
                c.width = Cm(widths[ci])
            if ri == 0:
                set_cell_shading(c, 'E9EEF4')
            elif ri % 2 == 0:
                set_cell_shading(c, 'F7F9FB')
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(8 if small else 9)
                    if ri == 0:
                        r.bold = True
    if rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement('w:tblHeader')
        tbl_header.set(qn('w:val'), 'true')
        tr_pr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_rich(p, raw):
    # Convert Markdown emphasis, code and link labels to readable Word runs.
    raw = re.sub(r'!\[([^]]*)\]\([^)]*\)', r'\1', raw)
    raw = re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', raw)
    bits = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', raw)
    for bit in bits:
        if not bit:
            continue
        if bit.startswith('**') and bit.endswith('**'):
            r = p.add_run(bit[2:-2]); r.bold = True
        elif bit.startswith('`') and bit.endswith('`'):
            r = p.add_run(bit[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            p.add_run(bit)


def add_image_pair(doc):
    doc.add_page_break()
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(7.2)
    table.columns[1].width = Cm(7.2)
    for ci, title in enumerate(('База, cell = 32', 'Фактор, cell = 64')):
        c = table.cell(0, ci)
        c.text = title
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(c, 'E9EEF4')
        set_cell_border(c)
        im = table.cell(1, ci)
        im.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        path = ROOT / 'artifacts/variant03/artifacts' / ('base_seed101.png' if ci == 0 else 'perturbed_seed101.png')
        im.paragraphs[0].add_run().add_picture(str(path), width=Cm(6.35), height=Cm(6.35))
        set_cell_border(im)
    doc.add_paragraph('Рисунок 1. Текстуры для seed 101 в одном масштабе.').style = 'Caption'


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.35)
sec.right_margin = Cm(2.0)
styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.15
styles['Normal'].paragraph_format.keep_together = True
for key, size, before, after in [('Title', 16, 0, 10), ('Heading 1', 12, 13, 5), ('Heading 2', 10.5, 9, 4)]:
    st = styles[key]
    st.font.name = 'Arial'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    ppr = st._element.pPr
    if ppr is not None:
        for border in ppr.findall(qn('w:pBdr')):
            ppr.remove(border)
styles['Caption'].font.name = 'Arial'
styles['Caption'].font.size = Pt(9)
styles['Caption'].font.italic = True
styles['Caption'].font.color.rgb = RGBColor(0,0,0)

lines = SOURCE.read_text().splitlines()
i = 0
in_code = False
while i < len(lines):
    line = lines[i].strip()
    if not line or line == '---':
        i += 1; continue
    if line.startswith('```'):
        in_code = not in_code; i += 1; continue
    if in_code:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(.4)
        r = p.add_run(line)
        r.font.name = 'Consolas'; r.font.size = Pt(8.5)
        i += 1; continue
    if line.startswith('|'):
        block = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i].strip()); i += 1
        if any('![Облачная' in b for b in block):
            add_image_pair(doc)
        elif 'Среднее база' in block[0] and 's_base' in block[0]:
            names = {'mean':'Средняя яркость', 'contrast':'Контраст', 'edge':'Граничная энергия', 'entropy':'Энтропия, бит/пиксель'}
            rows = [('Метрика','База, среднее','s базы','Фактор, среднее','s фактора','Δ','|Δ| / s базы','Итог')]
            for key, name in names.items():
                m=RESULT['metrics_table'][key]
                rows.append((name, f"{m['mean_base']:.4f}", f"{m['s_base']:.4f}", f"{m['mean_pert']:.4f}", f"{m['s_pert']:.4f}", f"{m['delta']:+.4f}", f"{m['delta_over_s']:.2f}", 'Да' if m['significant'] else 'Нет'))
            add_table(doc, rows, widths=[3.1,1.65,1.35,1.65,1.35,1.25,1.75,1.0], small=True)
        else:
            rows = [[re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', x.strip()) for x in b.strip('|').split('|')] for b in block]
            rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', x or '') for x in row)]
            if rows:
                add_table(doc, rows, small=len(rows[0]) > 4)
        continue
    if line.startswith('# '):
        # The title includes the actual topic and purpose in one Word Title paragraph.
        if line == '# Лабораторная работа № 1_1':
            title = doc.add_paragraph('Лабораторная работа 1 Подготовка среды и паспорт воспроизводимого AI эксперимента', style='Title')
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.paragraph_format.space_before = Cm(5)
        else:
            doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('## '):
        head = line[3:]
        if head.startswith('Подготовка рабочей среды'):
            pass
        else:
            doc.add_paragraph(head, style='Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], style='Heading 2')
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(.5)
        p.paragraph_format.first_line_indent = Cm(-.5)
        add_rich(p, line)
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); add_rich(p, line[2:])
    else:
        p = doc.add_paragraph(); add_rich(p, line.replace('  ', ' '))
        if line.startswith('**Папка работы:**'):
            doc.add_page_break()
    i += 1

doc.add_paragraph('Приложение А Полный паспорт эксперимента', style='Heading 1')
passport_rows = [
    ('Поле', 'Значение'),
    ('Цель и критерий', 'Проверить воспроизводимость генератора и влияние удвоения шага решётки. Эффект обнаружен, если |Δ| > 2 × s_base при n = 5.'),
    ('Данные и происхождение', 'Внешних данных нет. Изображения синтетические; исходные скрипты предоставлены с методическим архивом. Отдельная лицензия на код в архиве не указана.'),
    ('Модель', 'Учебный процедурный генератор lab01_generator.py, версия 1.0.'),
    ('Параметры', 'База: size 96, cell 32, octaves 4, persistence 0.5, quantize 8, interp smooth. Фактор: только cell 64. Серия seed 101–105, n = 5.'),
    ('Среда', 'macOS 26.0.1 arm64; CPython 3.12.12 в .venv; zlib 1.2.12. Сторонних пакетов нет; requirements.txt пуст.'),
    ('Команда', '.venv/bin/python src/lab01_experiment.py --variant 3 --out artifacts/variant03'),
    ('Результат', 'Для протокола seed 101, 101, 1101, 101: true, true, false, true. Уменьшение контраста, граничной энергии и энтропии обнаружено; изменение средней яркости не обнаружено.'),
    ('Ограничения', 'Пять seed, учебное правило 2s, одна модель и один набор параметров. Метрики не измеряют читаемость заголовка и художественное качество.'),
]
add_table(doc, passport_rows, widths=[3.5, 11.5])
doc.add_paragraph('Контрольные суммы исходного кода', style='Heading 2')
for name, sha in RESULT['script_sha256'].items():
    p = doc.add_paragraph(); p.add_run(name + ': ').bold = True; p.add_run(sha)

doc.add_paragraph('Приложение Б Полные отпечатки проверочных запусков', style='Heading 1')
rows = [('Запуск', 'Seed', 'SHA 256 полного массива пикселей')]
for idx, run in enumerate(RESULT['runs_repro'], 1):
    rows.append((str(idx), str(run['seed']), run['sha256_pixels']))
add_table(doc, rows, widths=[1.4, 1.5, 12.6], small=True)
doc.add_paragraph('Для запусков 1 и 2 SHA 256 файла PNG: ' + RESULT['runs_repro'][0]['sha256_png'] + '. Совпадение файла PNG установлено в этой среде; основной критерий — массив пикселей.')

doc.add_paragraph('Приложение В Диагностический журнал', style='Heading 1')
doc.add_paragraph('Дата основной серии: 24.09.2026 13:39:56 UTC. Дата фиксации гипотезы: 24.09.2026 13:39 UTC. Исполнитель: Власюк Данил Витальевич.')
rows = [('Проверка', 'Seed', 'SHA 256 пикселей', 'Наблюдение')]
for run in DIAG['runs']:
    rows.append((run['label'], str(run['seed']), run['sha256_pixels'], 'Предупреждение о незаданном seed' if run['warning'] else 'Повтор с явным seed'))
add_table(doc, rows, widths=[2.0, 1.5, 8.0, 4.0], small=True)
doc.add_paragraph('В двух запусках без --seed отпечатки различны; в двух запусках с --seed 101 совпадают. Причина выявленного расхождения — незаданное начальное состояние генератора.')
doc.add_paragraph('Журнал команд', style='Heading 2')
for command in [
    'python3.12 -m venv .venv',
    '.venv/bin/python -m pip freeze > requirements.txt',
    '.venv/bin/python src/lab01_generator.py --out artifacts/seed_diagnostic/без_seed_1',
    '.venv/bin/python src/lab01_generator.py --out artifacts/seed_diagnostic/без_seed_2',
    '.venv/bin/python src/lab01_generator.py --seed 101 --out artifacts/seed_diagnostic/seed_101_1',
    '.venv/bin/python src/lab01_generator.py --seed 101 --out artifacts/seed_diagnostic/seed_101_2',
    '.venv/bin/python src/lab01_experiment.py --variant 3 --out artifacts/variant03',
]:
    p = doc.add_paragraph(); r = p.add_run(command); r.font.name='Consolas'; r.font.size=Pt(8.5)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
