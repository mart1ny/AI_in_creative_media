from pathlib import Path
from format_markdown_report import format_report
import json,csv,tomllib,re,statistics
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];A=R/'artifacts';O=R/'reports'
e=tomllib.loads((A/'environment_and_results.toml').read_text());summ=list(csv.DictReader((A/'summary.csv').open()));per=list(csv.DictReader((A/'per_seed.csv').open()))
d=Document();s=d.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.5);s.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6)
for st in d.styles:
 for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
 if st.name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
  for key in list(fonts.attrib):
   if key.endswith('Theme'):del fonts.attrib[key]
  for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(22);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(13)
md=[]
def p(t,style=None):
 md.append(t+'\n');return d.add_paragraph(t,style)
def h(t,level=1):md.append('#'*(level+1)+' '+t+'\n');return d.add_heading(t,level)
def page(t):
 h(t);d.paragraphs[-1].paragraph_format.page_break_before=True
def table(headers,rows,widths):
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(str,row))+' |\n' for row in rows))
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right'):
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b);marg=OxmlElement('w:tcMar')
   for edge in ('top','bottom','left','right'):
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'75');x.set(qn('w:type'),'dxa');marg.append(x)
   pr.append(marg)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF2' if i==0 else ('F7F8FA' if i%2==0 else 'FFFFFF'));pr.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in pp.runs:run.font.size=Pt(9.5);run.bold=i==0
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def pic(name,caption,width=16.5):
 pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.add_run().add_picture(str(A/name),width=Cm(width));
 cap=p(caption);cap.paragraph_format.keep_with_next=False;cap.paragraph_format.space_after=Pt(10);cap.runs[0].font.size=Pt(10)
 md.append(f'![{caption}](../artifacts/{name})\n')
def code(txt):
 pp=p(txt);pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(8)
 for run in pp.runs:run.font.name='Courier New';run.font.size=Pt(9)

x=p('Искусственный интеллект в креативных технологиях');x.alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3):d.add_paragraph()
p('Лабораторная работа 3', 'Title').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Подготовка изображений и визуализация в Engee','Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Вариант 3', 'Heading 1').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Уменьшение изображения в четыре раза и возврат к исходному размеру').alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3):d.add_paragraph()
p('Выполнил Власюк Данил Витальевич\nГруппа Т26ИИКТ-МИК11').alignment=WD_ALIGN_PARAGRAPH.RIGHT
for _ in range(4):d.add_paragraph()
p('7 октября 2026 года').alignment=WD_ALIGN_PARAGRAPH.CENTER

page('1 Цель и методика эксперимента')
p('Цель работы — описать цифровое изображение, сравнить два способа получения яркости и измерить потери после изменения размера. Для варианта 3 выполнено уменьшение базовой сцены 128×128 до 32×32 и восстановление до 128×128. Сравнительный режим — уменьшение до 64×64. По пяти seed практический порог пересекла только энтропия; средний PSNR составил 35,24 дБ.')
h('Постановка и данные',2)
p('Вариант N=(1−1)×6+3=3: сцена 1 и фактор 3. Синтетическая сцена содержит диагональные градиенты каналов RGB, круг и гауссов шум. Параметры: n=128, lo=0.15, hi=0.85, noise=0.02, checker=0. Серия seed — 101, 102, 103, 104, 105. Все случайные значения берутся из локального MersenneTwister(seed). Изображения людей и внешние произведения не использовались.')
h('Гипотеза и допуск до запуска',2)
p('07.10.2026 в 20:23:23 MSK, до старта в 20:32:24 MSK, записано: средняя яркость изменится менее чем на 5%, контраст и энтропия снизятся, а PSNR при уменьшении вчетверо будет ниже, чем вдвое. Дополнительная серия при noise=0 проверяет, сохраняются ли потери без шума.')
p('Практический критерий по каждой метрике: |Δ|/|base|>0.05; по изображению: средний PSNR<30 дБ. 5% — учебный допуск относительного сдвига, а 30 дБ соответствует MSE=0.001 для [0,1]. Оба порога условны и не заменяют визуальную оценку. Учебный статистический критерий — |Δ|>2s, где s — выборочный разброс базовой метрики между seed; это эвристика методички без заданного уровня p.')
h('Среда и порядок расчёта',2)
p('Фактическая среда: Engee 26.9.2-H4; Julia 1.12.4; Linux x86_64; Images 0.26.2; Plots 1.41.6. Библиотека lab03_lib.jl от 19.09.2026 сохранена без изменений. Расчёт выполнен в /user/AI_in_creative_media/LR03, исходный вывод сохранён в ngscript и run_output.txt.')
code('include("lab03_lib.jl")\ng = luma_images(make_scene(101; scene_preset(1)...))\nh = apply_factor(g, 3)\nvariant_summary(3)')
p('Используются mean, std с corrected=true, энтропия 32 равных интервалов [0,1] и PSNR=10log10(1/MSE), MSE=mean((g−h)²). Сначала рассчитываются массивы Float64, затем отдельно сохраняются PNG N0f8 для просмотра. Поэтому экспортное квантование не входит в потери воздействия.')

page('2 Свойства сцены и воспроизводимость')
table(['Свойство','Значение'],[['Размер RGB и яркости','128×128'],['Тип цветного пикселя','RGB{Float64}'],['Диапазон каналов RGB','0.08115622264322489 … 0.9185814096363565'],['Объём RGB','393216 байт = 128×128×3×8'],['Объём яркости Float64','131072 байта = 128×128×8'],['Промежуточный размер варианта 3','32×32; в 16 раз меньше отсчётов'],['Возвращённый размер','128×128']],[6,10.5])
p('Ручная яркость вычислена как 0.299R+0.587G+0.114B, стандартная — Float64.(Gray.(img)). Максимум абсолютной разности равен 9.856187532797378×10⁻⁹; допуск 10⁻⁷ соблюдён. Это численное совпадение с допуском, а не побитовое равенство.')
p('SHA-256 стандартной яркости seed 101 при двух независимых построениях сцены совпал:')
code(e['sha_first'][:32]+'\n'+e['sha_first'][32:])
p('Хеш рассчитан по Float64 в порядке vec Julia по столбцам, little-endian. Полная 64-символьная сумма записана в паспорте; исходный массив сохранён как before_seed101.f64. Контрольная сумма PNG отличается, поскольку у него другая структура и точность данных.')
pic('comparison_scene.png','Рисунок 1 — сцена seed 101 и два режима обработки')
pic('circle_edge_detail.png','Рисунок 2 — фрагмент границы круга, строки 27–58 и столбцы 45–76',width=14)


page('3 Результаты варианта 3')
p('Средние рассчитаны по пяти seed. s характеризует разброс базовой метрики между seed, а контраст — std по пикселям одного изображения. Эти две величины имеют разный смысл.')
rows=[]
labels={'mean':'Среднее','contrast':'Контраст','entropy':'Энтропия бит'}
for z in summ:
 if int(z['variant'])!=3:continue
 rows.append([labels[z['metric']],f"{float(z['base']):.6f}",f"{float(z['s']):.3e}",f"{float(z['after']):.6f}",f"{float(z['delta']):.6f}",f"{float(z['ratio']):.3f}",'Да' if z['stat_significant']=='true' else 'Нет'])
table(['Метрика','База','s','После','Δ','|Δ|/s','>2s'],rows,[2.4,2.2,2.3,2.2,2.2,1.9,1.3])
pic('histograms_comparison.png','Рисунок 3 — гистограммы исходных Float64, 32 интервала, seed 101')
h('Оценка каждой метрики',2)
p('Средняя яркость снизилась с 0.471942 до 0.471452: Δ=−0.000490, относительное изменение 0.104%. Критерий 2s выполнен (5.842s), практический допуск 5% соблюдён. Значительного общего затемнения по этому допуску нет.')
p('Контраст снизился с 0.064246 до 0.061669: Δ=−0.002577, изменение 4.011%. Критерий 2s выполнен (16.737s), практический допуск 5% соблюдён. Потеря контраста есть, но меньше заранее выбранного порога.')
p('Энтропия снизилась с 2.687328 до 2.486132 бит: Δ=−0.201196 бит, изменение 7.487%. Критерий 2s выполнен (45.277s) и практический порог 5% превышен. На гистограмме шумовые значения концентрируются в более узких группах.')
p('Средний PSNR пяти seed равен 35.244413 дБ, выборочный разброс 0.073074 дБ; диапазон 35.141812…35.298577 дБ. Все значения выше 30 дБ. По принятому порогу суммарная MSE остаётся допустимой, хотя граница круга визуально сглажена.')

page('4 Сравнение масштабов и контроль шума')
table(['Режим','Среднее','Контраст','Энтропия бит','PSNR дБ'],[['Исходный','0.471942','0.064246','2.687328','—'],['128 → 64 → 128','0.471942','0.061871','2.437457','36.874023'],['128 → 32 → 128','0.471452','0.061669','2.486132','35.244413']],[4.4,2.5,2.7,3.5,3.4])
p('При уменьшении вдвое средняя яркость сохраняется с точностью округления; критерий 2s для неё не выполнен. Контраст снижается на 3.697% (15.425s), энтропия — на 9.298% (56.230s). Практически изменяется только энтропия. Средний PSNR выше результата ×4 на 1.629610 дБ.')
p('Более сильное уменьшение снижает PSNR и контраст, однако энтропия после ×4 выше, чем после ×2, на 0.048675 бит. Энтропия фиксированных интервалов описывает распределение яркости, а не пространственное качество: интерполяция границ создаёт промежуточные значения. Поэтому энтропию нельзя ранжировать как универсальную меру сохранённых деталей.')
h('Пять запусков основного воздействия',2)
r=[x for x in per if float(x['noise'])==.02 and int(x['factor'])==3]
table(['Seed','Среднее после','Контраст после','Энтропия после','MSE','PSNR дБ'],[[x['seed'],f"{float(x['mean_after']):.6f}",f"{float(x['contrast_after']):.6f}",f"{float(x['entropy_after']):.6f}",f"{float(x['mse']):.3e}",f"{float(x['psnr']):.6f}"] for x in r],[1.3,3,3,3,3,3.2])
h('Контроль при отсутствии шума',2)
table(['Шум','Режим','Средняя MSE','Средний PSNR дБ'],[['0.02','×2','0.000205413','36.874023'],['0.02','×4','0.000298957','35.244413'],['0','×2','0.000057993','42.366227'],['0','×4','0.000114540','39.410417']],[2.3,2.4,5.9,5.9])
p('Без шума ошибка снижается, но не исчезает; ×4 по-прежнему даёт худший PSNR. Это подтверждает вклад структуры сцены, в том числе резкой границы круга, а не только сглаживания шума. При noise=0 все seed дают одинаковые данные и s=0: критерий |Δ|/s здесь не применяется. Эта серия служит диагностикой, а не пятью независимыми наблюдениями.')

page('5 Анализ ограничений и вывод')
h('Чувствительность к практическому допуску',2)
p('При исходном пороге 5% практически значима только энтропия. Для проверки чувствительности, без изменения основного решения, при пороге 3% значимы контраст и энтропия; при 10% ни одна из трёх метрик его не пересекает. Средняя яркость сдвинулась лишь на 0.104%. При PSNR-пороге 30 дБ оба режима допустимы; при более строгом 36 дБ режим ×4 пересёк бы порог, а ×2 — нет. Поэтому порог должен следовать из задачи и фиксироваться до расчёта.')
h('Границы интерпретации',2)
p('Критерий 2s сравнивает сдвиг с вариативностью базовой сцены. Он не учитывает распределение парных разностей, не задаёт уровень ошибки первого рода и не является формальным тестом с p<0.05. В отчёте «статистически значимо» означает только выполнение учебного критерия. Пять seed меняют реализацию шума одной сцены, а не представляют выборку из разных изображений.')
p('Вывод ограничен синтетической сценой и штатной imresize с параметрами по умолчанию в данной версии Images. Другие изображения, интерполяция, шум или границы изменят результат. Глобальные среднее, std и энтропия могут скрыть локальную потерю деталей; PSNR измеряет численное отличие и не оценивает красоту или читаемость текста. Фиксированная разбивка гистограммы влияет на энтропию.')
p('Контроль без шума показал ненулевые потери на структуре сцены. Более сильное уменьшение не обеспечило монотонного изменения энтропии: она зависит от перераспределения значений по интервалам. Для следующей проверки можно взять собственную фотографию тестовой таблицы с тонкими линиями и текстом, подтвердить авторство, сохранить оригинал и повторить оба масштаба при тех же критериях; отдельно оценить читаемость деталей.')
h('Проверки выполнения',2)
p('Подтверждены диапазон [0,1], размер 32×32 после уменьшения, возврат 128×128, совпадение двух хешей, численное совпадение способов яркости и воспроизведение демо. Независимый verify_results.py пересчитал сводку из CSV, хеши Float64 и PSNR по сохранённым массивам; RESULT PASS. Библиотека преподавателя совпала с исходником по байтам. Исполненный ngscript содержит сохранённый stdout и не содержит ошибок.')
h('Вывод',2)
p('Предварительная гипотеза подтверждена: средняя яркость меняется слабо, контраст и энтропия снижаются, а уменьшение вчетверо даёт меньший PSNR, чем вдвое. В варианте 3 все три метрики пересекают учебный порог 2s, но только снижение энтропии на 7.487% превышает практический порог 5%. Средний PSNR 35.244413 дБ остаётся выше 30 дБ. Возврат исходного размера не восстанавливает утраченные детали; оценку качества следует дополнять осмотром границ и деталей.')

answers=(O/'control_answers.md').read_text().split('\n\n')[1:]
for idx,start in enumerate((0,5)):
 page('6 Контрольные вопросы' if idx==0 else '6 Контрольные вопросы продолжение')
 for txt in answers[start:start+5]:
  txt=re.sub(r'\*\*(.*?)\*\*',r'\1',txt)
  label,body=txt.split('. ',1);title,body=body.split('. ',1)
  h(label+' '+title,2);p(body)

page('7 Паспорт и комплект для сдачи')
p('Студент — Власюк Данил Витальевич, Т26ИИКТ-МИК11. Вариант 3. Дата выполнения — 07.10.2026, Europe/Moscow. Предварительная запись — 20:23:23 MSK; расчёт — 20:32:24…20:32:44 MSK. Серверное время вывода на три часа меньше московского.')
p('Среда — Engee 26.9.2-H4, Julia 1.12.4, Linux x86_64-linux-gnu. Images 0.26.2, Plots 1.41.6; Random, Statistics, SHA, TOML, Tar, Dates, InteractiveUtils — стандартные библиотеки Julia. Автоматические метаданные ngscript указывают IJulia 1.9.3; фактическая версия 1.12.4 подтверждена VERSION и versioninfo в сохранённом stdout.')
p('Исходная библиотека lab03_lib.jl от 19.09.2026 — 5027 байт, без изменений. SHA-256:')
code(e['library_sha256'][:32]+'\n'+e['library_sha256'][32:])
p('Параметры сцены: n=128, lo=0.15, hi=0.85, noise=0.02, checker=0. Seed 101…105, локальный MersenneTwister. Основная цепочка 128 → 32 → 128; сравнительная 128 → 64 → 128. Гистограмма — 32 равных интервала [0,1], индекс min(32,max(1,ceil(Int,v*32))). Яркость Gray Float64; статистический критерий |Δ|>2s; практические пороги 5% и 30 дБ.')
p('Контрольная сумма яркости seed 101 и повтор равны. Полный хеш приведён на странице результатов воспроизводимости и в passport.json. Байтовая схема — Float64 IEEE754, little-endian, vec Julia по столбцам. Изображения сохранены в N0f8 только для просмотра; метрики рассчитаны по Float64. Демонстрация воспроизведена. Итоговая проверка — RESULT PASS.')
h('Архив артефактов',2)
p('ЛР03_Вариант_3_Сдача.zip лежит в папке LR3 рядом с README.md. Облачная папка — /user/AI_in_creative_media/LR03. Сдаваемый комплект содержит этот отчёт, исполняемый LR03_3.ngscript с stdout, LR03_3.jl и lab03_lib.jl, PNG до и после, гистограммы, CSV по seed и сводку, исходные Float64, паспорт, гипотезу до запуска и журнал проверки. Полный паспорт доступен в reports/passport.md и reports/passport.json.')
code('В Engee из папки с исходниками:\ninclude("LR03_3.jl")\n\nНезависимая проверка локально из LR3:\npython3 scripts/verify_results.py')
h('Источники',2)
p('1. Методические указания «Лабораторная работа № 3. Подготовка изображений и визуализация в Engee», предоставленный DOCX; разделы 2–9.\n2. Архив преподавателя: lab03_lib.jl, variants_lr03.json, журнал демо от 19.09.2026.\n3. Собственные результаты Engee от 07.10.2026: per_seed.csv, summary.csv, environment_and_results.toml, versioninfo.txt и сохранённый вывод LR03_3.ngscript.')

# Номера страниц полезны для навигации по отчёту.
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
d.core_properties.author='Власюк Данил Витальевич';d.core_properties.title='Лабораторная работа 3 Подготовка изображений и визуализация в Engee Вариант 3'
d.save(O/'ЛР03_Вариант_3_Отчёт.docx')
(O/'ЛР03_Вариант_3_Отчёт.md').write_text(format_report('\n'.join(md)))
print(O/'ЛР03_Вариант_3_Отчёт.docx')
